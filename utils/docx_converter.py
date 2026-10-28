"""
DOCX -> structured, LLM-safe JSON.

Design goals (to avoid hallucination downstream):
  * Document order is preserved: every block has an index.
  * Tables keep row/column relationships: each data cell is addressable by its
    column header, merged cells are resolved explicitly, nested tables are kept.
  * Tracked changes, comments and footnotes are labelled, never blended silently
    into body text. Paragraph "text" is the ACCEPTED view (deletions removed).
  * Things we cannot interpret (images, etc.) are emitted as explicit placeholders.
"""
import io
import re
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}
_TXBX = etree.XPath("./descendant::w:txbxContent[not(ancestor::mc:Fallback)]", namespaces=NS)
_DOCPR = etree.XPath(".//wp:docPr", namespaces=NS)
_IN_FALLBACK = etree.XPath("boolean(ancestor::mc:Fallback)", namespaces=NS)


class _Ctx:
    def __init__(self, doc):
        self.doc = doc
        self.stats = {"paragraphs": 0, "tables": 0, "images": 0,
                      "tracked_changes": 0, "comments": 0, "nested_tables": 0}
        self._num_cache: dict = {}

    def list_kind(self, num_id: str, ilvl: str) -> str:
        key = (num_id, ilvl)
        if key in self._num_cache:
            return self._num_cache[key]
        kind = "unknown"
        try:
            root = self.doc.part.numbering_part.element
            num = root.xpath(f"./w:num[@w:numId='{num_id}']")[0]
            abs_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
            lvl = root.xpath(
                f"./w:abstractNum[@w:abstractNumId='{abs_id}']/w:lvl[@w:ilvl='{ilvl}']")[0]
            fmt = lvl.find(qn("w:numFmt")).get(qn("w:val"))
            kind = "bullet" if fmt == "bullet" else "numbered"
        except Exception:
            pass
        self._num_cache[key] = kind
        return kind


# ---------------------------------------------------------------- paragraphs
def _run_text(r) -> str:
    out = []
    for c in r:
        t = c.tag
        if t in (qn("w:t"), qn("w:delText")):
            out.append(c.text or "")
        elif t == qn("w:tab"):
            out.append("\t")
        elif t in (qn("w:br"), qn("w:cr")):
            out.append("\n")
        elif t == qn("w:noBreakHyphen"):
            out.append("-")
        elif t == qn("w:footnoteReference"):
            out.append(f"[^fn{c.get(qn('w:id'))}]")
    return "".join(out)


def _para_text(p_el):
    """Return (accepted_text, tracked_changes list)."""
    accepted: list[str] = []
    tracked: list[dict] = []

    def add_tracked(kind, text):
        if not text:
            return
        if tracked and tracked[-1]["type"] == kind:
            tracked[-1]["text"] += text
        else:
            tracked.append({"type": kind, "text": text})

    def walk(el, state):
        for ch in el:
            tag = ch.tag
            if tag in (qn("w:ins"), qn("w:moveTo")):
                walk(ch, "inserted")
            elif tag in (qn("w:del"), qn("w:moveFrom")):
                walk(ch, "deleted")
            elif tag == qn("w:r"):
                text = _run_text(ch)
                if state == "deleted":
                    add_tracked("deleted", text)
                else:
                    accepted.append(text)
                    if state == "inserted":
                        add_tracked("inserted", text)
            elif tag in (qn("w:hyperlink"), qn("w:smartTag"), qn("w:fldSimple"),
                         qn("w:sdt"), qn("w:sdtContent")):
                walk(ch, state)

    walk(p_el, None)
    return "".join(accepted), tracked


def _parse_paragraph(p_el, ctx: _Ctx, idx_holder) -> list[dict]:
    text, tracked = _para_text(p_el)
    blocks: list[dict] = []

    # style / heading
    style_name = ""
    ppr = p_el.find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:pStyle")) is not None:
        sid = ppr.find(qn("w:pStyle")).get(qn("w:val"))
        try:
            style_name = ctx.doc.styles.get_by_id(sid, 1).name  # 1 = paragraph style
        except Exception:
            style_name = sid or ""

    block: dict[str, Any] = {"type": "paragraph"}
    m = re.match(r"(?i)heading\s*(\d)", style_name or "")
    if m:
        block = {"type": "heading", "level": int(m.group(1))}
    elif (style_name or "").lower() == "title":
        block = {"type": "heading", "level": 0}

    if ppr is not None and ppr.find(qn("w:numPr")) is not None:
        numpr = ppr.find(qn("w:numPr"))
        num_id = numpr.find(qn("w:numId"))
        ilvl = numpr.find(qn("w:ilvl"))
        nid = num_id.get(qn("w:val")) if num_id is not None else None
        lvl = ilvl.get(qn("w:val")) if ilvl is not None else "0"
        if nid and nid != "0":
            block["list"] = {"kind": ctx.list_kind(nid, lvl), "level": int(lvl)}

    if "list" not in block:  # numbering inherited from the paragraph style
        sl = (style_name or "").lower()
        if sl.startswith("list bullet"):
            block["list"] = {"kind": "bullet", "level": 0}
        elif sl.startswith("list number"):
            block["list"] = {"kind": "numbered", "level": 0}

    text = text.strip("\n")
    if text.strip():
        block["text"] = text
        ctx.stats["paragraphs"] += 1
        if tracked:
            block["tracked_changes"] = tracked
            ctx.stats["tracked_changes"] += len(tracked)
        refs = [c.get(qn("w:id")) for c in p_el.iter(qn("w:commentReference"))]
        if refs:
            block["comment_ids"] = refs
        blocks.append(block)
    elif tracked:  # paragraph entirely deleted
        ctx.stats["tracked_changes"] += len(tracked)
        blocks.append({"type": "paragraph", "text": "", "tracked_changes": tracked})

    # images / drawings
    for tag in ("w:drawing", "w:pict"):
        for d in p_el.iter(qn(tag)):
            if _IN_FALLBACK(d):
                continue
            props = _DOCPR(d)
            descr = ""
            if props:
                descr = props[0].get("descr") or props[0].get("name") or ""
            ctx.stats["images"] += 1
            blocks.append({"type": "image_placeholder",
                           "alt_text": descr,
                           "note": "Image/drawing present in source; content NOT interpreted."})

    # text boxes
    for tb in _TXBX(p_el):
        inner = _parse_blocks(tb, ctx, idx_holder, assign_index=False)
        if inner:
            blocks.append({"type": "text_box", "content": inner})
    return blocks


# -------------------------------------------------------------------- tables
def _cell_info(tc, ctx):
    tcpr = tc.find(qn("w:tcPr"))
    span, vm = 1, None
    if tcpr is not None:
        gs = tcpr.find(qn("w:gridSpan"))
        if gs is not None:
            span = int(gs.get(qn("w:val")))
        v = tcpr.find(qn("w:vMerge"))
        if v is not None:
            vm = v.get(qn("w:val")) or "continue"
    inner = _parse_blocks(tc, ctx, [0], assign_index=False)
    texts, nested = [], []
    for b in inner:
        if b["type"] in ("paragraph", "heading") and b.get("text"):
            prefix = ""
            if "list" in b:
                prefix = "- " if b["list"]["kind"] != "numbered" else "* "
            texts.append(prefix + b["text"])
        elif b["type"] == "table":
            nested.append(b)
            ctx.stats["nested_tables"] += 1
        elif b["type"] == "image_placeholder":
            texts.append("[image]")
    return {"text": "\n".join(texts), "span": span, "vmerge": vm, "nested": nested}


def _parse_table(tbl_el, ctx: _Ctx) -> dict:
    ctx.stats["tables"] += 1
    rows_raw, flagged_header = [], 0
    for tr in tbl_el.findall(qn("w:tr")):
        trpr = tr.find(qn("w:trPr"))
        if trpr is not None and trpr.find(qn("w:tblHeader")) is not None \
                and flagged_header == len(rows_raw):
            flagged_header += 1
        rows_raw.append([_cell_info(tc, ctx) for tc in tr.findall(qn("w:tc"))])

    # resolve into a rectangular grid (merged cells resolved explicitly)
    grid: list[list[dict]] = []
    for r, cells in enumerate(rows_raw):
        row: list[dict] = []
        for ci in cells:
            c0 = len(row)
            for k in range(ci["span"]):
                col = c0 + k
                if ci["vmerge"] == "continue" and r > 0 and col < len(grid[r - 1]):
                    origin = grid[r - 1][col]["origin"]
                    text = grid[origin[0]][origin[1]]["text"]
                    nested = []
                else:
                    origin = (r, c0)
                    text = ci["text"]
                    nested = ci["nested"] if k == 0 else []
                row.append({"text": text, "origin": origin,
                            "span_part": k > 0, "nested": nested})
        grid.append(row)
    ncols = max((len(r) for r in grid), default=0)
    for r, row in enumerate(grid):
        while len(row) < ncols:
            row.append({"text": "", "origin": (r, len(row)), "span_part": False, "nested": []})

    # merged-cell descriptions
    merged = []
    for r, row in enumerate(grid):
        for c, cell in enumerate(row):
            if cell["origin"] == (r, c) and not cell["span_part"]:
                colspan = 1
                while c + colspan < ncols and grid[r][c + colspan]["span_part"] \
                        and grid[r][c + colspan]["origin"] == (r, c):
                    colspan += 1
                rowspan = 1
                while r + rowspan < len(grid) and grid[r + rowspan][c]["origin"] == (r, c):
                    rowspan += 1
                if colspan > 1 or rowspan > 1:
                    merged.append({"row": r, "col": c, "rowspan": rowspan,
                                   "colspan": colspan, "text": cell["text"]})

    header_inferred = flagged_header == 0
    n_head = flagged_header or (1 if grid else 0)
    if header_inferred and len(grid) > 2 and any(
            m["row"] == 0 and m["colspan"] > 1 for m in merged):
        n_head = 2  # grouped header: merged top row + column-name row
    headers: list[str] = []
    for c in range(ncols):
        parts = []
        for r in range(n_head):
            t = (grid[r][c]["text"] or "").strip()
            if t and t not in parts:
                parts.append(t)
        headers.append(" / ".join(parts))

    text_grid = [[(cell["text"] or "") for cell in row] for row in grid]
    unique = all(headers) and len(set(headers)) == len(headers)
    records = None
    if unique and len(grid) > n_head:
        records = [{headers[c]: text_grid[r][c] for c in range(ncols)}
                   for r in range(n_head, len(grid))]

    nested = {f"{r},{c}": cell["nested"] for r, row in enumerate(grid)
              for c, cell in enumerate(row) if cell["nested"]}

    table = {"type": "table", "num_rows": len(grid), "num_cols": ncols,
             "header_rows": n_head, "header_inferred": header_inferred,
             "headers": headers, "records": records, "cells": text_grid,
             "merged_cells": merged}
    if nested:
        table["nested_tables"] = nested
    if records is None and grid:
        table["note"] = ("Headers are empty or duplicated, so 'records' is omitted; "
                         "use 'cells' with 'headers' by column index.")
    return table


# -------------------------------------------------------------------- blocks
def _parse_blocks(container, ctx: _Ctx, idx_holder, assign_index=True) -> list[dict]:
    out: list[dict] = []
    for ch in container:
        if ch.tag == qn("w:p"):
            new = _parse_paragraph(ch, ctx, idx_holder)
        elif ch.tag == qn("w:tbl"):
            new = [_parse_table(ch, ctx)]
        elif ch.tag in (qn("w:sdt"), qn("w:sdtContent")):
            content = ch.find(qn("w:sdtContent")) if ch.tag == qn("w:sdt") else ch
            new = _parse_blocks(content, ctx, idx_holder, assign_index=False) \
                if content is not None else []
        else:
            continue
        for b in new:
            if assign_index:
                b["index"] = idx_holder[0]
                idx_holder[0] += 1
            out.append(b)
    return out


# ------------------------------------------------------------------ markdown
def _esc(s: str) -> str:
    return (s or "").replace("|", "\\|").replace("\n", "<br>")


def _table_md(t: dict) -> str:
    if not t["num_cols"]:
        return ""
    heads = [h or " " for h in t["headers"]]
    lines = ["| " + " | ".join(_esc(h) for h in heads) + " |",
             "|" + "|".join("---" for _ in heads) + "|"]
    for r in range(t["header_rows"], t["num_rows"]):
        cells = []
        for c, v in enumerate(t["cells"][r]):
            if f"{r},{c}" in t.get("nested_tables", {}):
                v += " [nested table: see JSON]"
            cells.append(_esc(v))
        lines.append("| " + " | ".join(cells) + " |")
    if t["merged_cells"]:
        lines.append("")
        lines.append("_Merged cells (value repeated across merged area above): " + "; ".join(
            f"r{m['row']}c{m['col']} spans {m['rowspan']}x{m['colspan']}"
            for m in t["merged_cells"]) + "_")
    return "\n".join(lines)


def _to_markdown(blocks: list[dict]) -> str:
    parts = []
    for b in blocks:
        t = b["type"]
        if t == "heading":
            parts.append("#" * max(1, b["level"] or 1) + " " + b["text"])
        elif t == "paragraph":
            txt = b.get("text", "")
            if "list" in b:
                pad = "  " * b["list"]["level"]
                mark = "1." if b["list"]["kind"] == "numbered" else "-"
                parts.append(f"{pad}{mark} {txt}")
            elif txt:
                parts.append(txt)
            if b.get("tracked_changes"):
                parts.append("  _(tracked: " + "; ".join(
                    f"{c['type']} '{c['text']}'" for c in b["tracked_changes"]) + ")_")
        elif t == "table":
            parts.append(_table_md(b))
        elif t == "image_placeholder":
            parts.append(f"[IMAGE PLACEHOLDER: {b['alt_text'] or 'no alt text'} — not interpreted]")
        elif t == "text_box":
            parts.append("> [text box]\n" + "\n".join(
                "> " + ln for ln in _to_markdown(b["content"]).splitlines()))
    # keep consecutive list items tight, blank line elsewhere
    md, prev_list = [], False
    for p in parts:
        is_list = bool(re.match(r"\s*(-|1\.) ", p))
        md.append(("\n" if not (is_list and prev_list) and md else "") + p)
        prev_list = is_list
    return "\n".join(md).strip() + "\n"


# --------------------------------------------------------------- side parts
def _side_part(doc, suffix):
    for part in doc.part.package.iter_parts():
        if str(part.partname).endswith(suffix):
            return etree.fromstring(part.blob)
    return None


def _comments(doc, ctx):
    root = _side_part(doc, "/word/comments.xml")
    out = []
    if root is not None:
        for c in root.findall(qn("w:comment")):
            text = "\n".join(_para_text(p)[0] for p in c.iter(qn("w:p"))).strip()
            out.append({"id": c.get(qn("w:id")), "author": c.get(qn("w:author")),
                        "date": c.get(qn("w:date")), "text": text})
    ctx.stats["comments"] = len(out)
    return out


def _footnotes(doc):
    root = _side_part(doc, "/word/footnotes.xml")
    out = []
    if root is not None:
        for f in root.findall(qn("w:footnote")):
            if f.get(qn("w:type")) in ("separator", "continuationSeparator"):
                continue
            text = "\n".join(_para_text(p)[0] for p in f.iter(qn("w:p"))).strip()
            if text:
                out.append({"ref": f"[^fn{f.get(qn('w:id'))}]", "text": text})
    return out


def _headers_footers(doc):
    seen, out = set(), []
    for i, s in enumerate(doc.sections):
        for kind, part in (("header", s.header), ("footer", s.footer)):
            txt = "\n".join(p.text for p in part.paragraphs if p.text.strip())
            if txt and (kind, txt) not in seen:
                seen.add((kind, txt))
                out.append({"section": i, "kind": kind, "text": txt})
    return out


# -------------------------------------------------------------------- public
def convert_docx(data: bytes, filename: str = "") -> dict:
    doc = Document(io.BytesIO(data))
    ctx = _Ctx(doc)
    blocks = _parse_blocks(doc.element.body, ctx, [0])
    comments = _comments(doc, ctx)
    footnotes = _footnotes(doc)
    warnings = []
    if ctx.stats["images"]:
        warnings.append(f"{ctx.stats['images']} image(s)/drawing(s) were not interpreted; "
                        "see image_placeholder blocks.")
    if ctx.stats["tracked_changes"]:
        warnings.append("Document contains tracked changes. 'text' shows the accepted view; "
                        "see 'tracked_changes' per paragraph for insertions/deletions.")
    return {
        "filename": filename,
        "format_notes": ("Blocks are in document order. Table cells are resolved: merged cells "
                         "repeat their value across the merged area; see table.merged_cells. "
                         "Use table.records (header->value) to read rows."),
        "blocks": blocks,
        "markdown": _to_markdown(blocks),
        "comments": comments,
        "footnotes": footnotes,
        "headers_footers": _headers_footers(doc),
        "stats": ctx.stats,
        "warnings": warnings,
    }
